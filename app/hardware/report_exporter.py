from html import escape
from pathlib import Path

from PySide6.QtCore import QSizeF
from PySide6.QtGui import QPageSize, QTextDocument
from PySide6.QtPrintSupport import QPrinter, QPrinterInfo


class ReportExporter:
    _THERMAL_PRINTER_MARKERS = (
        "thermal",
        "receipt",
        "xprinter",
        "bixolon",
        "epson tm",
        "star tsp",
        "zjiang",
        "rongta",
        "80mm",
        "58mm",
        "pos-80",
        "pos 80",
        "pos-58",
        "pos 58",
    )

    def print_report(self, report_data: dict) -> str:
        available = [
            printer
            for printer in QPrinterInfo.availablePrinters()
            if not self.is_thermal_printer(printer.printerName())
        ]
        if not available:
            raise OSError("Tidak ada printer dokumen non-thermal yang tersedia.")

        default = QPrinterInfo.defaultPrinter()
        selected = next(
            (
                printer
                for printer in available
                if not default.isNull()
                and printer.printerName() == default.printerName()
            ),
            available[0],
        )
        printer = QPrinter(selected, QPrinter.PrinterMode.HighResolution)
        printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        self._create_document(report_data).print_(printer)
        if printer.printerState() == QPrinter.PrinterState.Error:
            raise OSError(
                f"Printer {selected.printerName()} melaporkan kesalahan."
            )
        return selected.printerName()

    def save_pdf(self, report_data: dict, file_path: str | Path) -> Path:
        output_path = Path(file_path)
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(str(output_path))
        printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        self._create_document(report_data).print_(printer)
        if printer.printerState() == QPrinter.PrinterState.Error:
            raise OSError("PDF laporan gagal dibuat.")
        if not output_path.is_file() or output_path.stat().st_size == 0:
            raise OSError("PDF laporan tidak ditemukan atau kosong.")
        return output_path

    def save_docx(self, report_data: dict, file_path: str | Path) -> Path:
        from docx import Document
        from docx.enum.section import WD_ORIENT

        output_path = Path(file_path)
        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = (
            section.page_height,
            section.page_width,
        )
        document.add_heading(report_data["title"], level=0)
        document.add_paragraph(f"Periode: {report_data['period']}")
        document.add_paragraph(f"Dibuat: {report_data['generated_at']}")
        document.add_paragraph(report_data["summary"])
        if report_data.get("warning"):
            document.add_paragraph(report_data["warning"])

        table = document.add_table(rows=1, cols=len(report_data["headers"]))
        table.style = "Table Grid"
        for cell, heading in zip(table.rows[0].cells, report_data["headers"]):
            cell.text = str(heading)
        for row in report_data["rows"]:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)

        document.save(output_path)
        if not output_path.is_file() or output_path.stat().st_size == 0:
            raise OSError("DOCX laporan tidak ditemukan atau kosong.")
        return output_path

    @classmethod
    def is_thermal_printer(cls, name: str) -> bool:
        normalized = name.casefold()
        return any(marker in normalized for marker in cls._THERMAL_PRINTER_MARKERS)

    def _create_document(self, report_data: dict) -> QTextDocument:
        document = QTextDocument()
        document.setDocumentMargin(36)
        document.setHtml(self.format_report_html(report_data))
        return document

    @staticmethod
    def format_report_html(report_data: dict) -> str:
        headers = "".join(
            f"<th>{escape(str(header))}</th>"
            for header in report_data["headers"]
        )
        rows = "".join(
            "<tr>"
            + "".join(f"<td>{escape(str(value))}</td>" for value in row)
            + "</tr>"
            for row in report_data["rows"]
        )
        warning = report_data.get("warning", "")
        warning_html = (
            f"<p class='warning'>{escape(str(warning))}</p>" if warning else ""
        )
        return f"""
        <html><head><style>
          body {{ font-family: Arial, sans-serif; font-size: 9pt; color: #172554; }}
          h1 {{ font-size: 18pt; }}
          p {{ margin: 4px 0; }}
          table {{ width: 100%; border-collapse: collapse; margin-top: 14px; }}
          th {{ background: #e2e8f0; text-align: left; }}
          th, td {{ border: 1px solid #cbd5e1; padding: 5px; }}
          .warning {{ color: #9a3412; font-weight: bold; }}
        </style></head><body>
          <h1>{escape(str(report_data['title']))}</h1>
          <p>Periode: {escape(str(report_data['period']))}</p>
          <p>Dibuat: {escape(str(report_data['generated_at']))}</p>
          <p>{escape(str(report_data['summary']))}</p>
          {warning_html}
          <table><thead><tr>{headers}</tr></thead><tbody>{rows}</tbody></table>
        </body></html>
        """
