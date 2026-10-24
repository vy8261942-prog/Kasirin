import csv
import sqlite3
from pathlib import Path
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.core.exceptions import AuthorizationError, ValidationError
from app.models import Category, Product, StockMovement, Supplier, Unit, User
from app.services.master_data_service import MasterDataService


PRODUCT_HEADERS = (
    "SKU",
    "Barcode",
    "Nama Produk",
    "Kategori",
    "Satuan",
    "Simbol Satuan",
    "Kode Supplier",
    "Harga Modal",
    "Harga Jual",
    "Stok Awal",
    "Batas Minimum Stok",
    "Aktif",
)


class ProductImportError(ValueError):
    pass


class ProductImportService:
    def __init__(self, session: Session):
        self.session = session
        self.master_data = MasterDataService(session)

    def preview(self, file_path: str | Path) -> list[dict[str, object]]:
        path = Path(file_path)
        suffix = path.suffix.casefold()
        if suffix == ".csv":
            return self._read_csv(path)
        if suffix == ".xlsx":
            return self._read_excel(path)
        raise ProductImportError("Format file harus .xlsx atau .csv.")

    def import_file(self, actor: User, file_path: str | Path) -> dict:
        return self.import_records(actor, self.preview(file_path))

    def import_database(self, actor: User, file_path: str | Path) -> dict:
        records = self._read_database(Path(file_path))
        return self.import_records(
            actor,
            records,
            create_missing_units=True,
            allow_missing_suppliers=True,
        )

    def preview_database(self, file_path: str | Path) -> list[dict[str, object]]:
        return self._read_database(Path(file_path))

    def import_records(
        self,
        actor: User,
        records: list[dict[str, object]],
        *,
        create_missing_units: bool = False,
        allow_missing_suppliers: bool = False,
    ) -> dict:
        if actor.role.name.upper() != "OWNER":
            raise AuthorizationError("Hanya OWNER yang dapat mengimpor produk.")
        if not records:
            raise ProductImportError("File tidak berisi baris produk.")

        errors: list[str] = []
        seen_skus: set[str] = set()
        created = 0
        updated = 0
        for row_number, record in enumerate(records, start=2):
            try:
                normalized = self._normalize_record(record)
                sku_key = str(normalized["sku"]).casefold()
                if sku_key in seen_skus:
                    raise ProductImportError("SKU duplikat di dalam file.")
                seen_skus.add(sku_key)
                self._validate_references(
                    normalized,
                    create_missing_units=create_missing_units,
                    allow_missing_suppliers=allow_missing_suppliers,
                )
            except (ProductImportError, TypeError, ValueError) as exc:
                errors.append(f"Baris {row_number}: {exc}")

        if errors:
            return {"created": 0, "updated": 0, "errors": errors}

        try:
            for row_number, record in enumerate(records, start=2):
                normalized = self._normalize_record(record)
                category = self._find_category(str(normalized["category"]))
                unit = self._find_unit(
                    str(normalized["unit"]),
                    str(normalized["unit_symbol"]),
                )
                if category is None:
                    category = Category(name=str(normalized["category"]).strip())
                    self.session.add(category)
                    self.session.flush()
                elif category.deleted_at is not None:
                    category.deleted_at = None
                if unit is None and create_missing_units:
                    unit = Unit(
                        name=str(normalized["unit"]).strip(),
                        symbol=str(normalized["unit_symbol"]).strip(),
                    )
                    self.session.add(unit)
                    self.session.flush()
                elif unit is not None and unit.deleted_at is not None:
                    unit.deleted_at = None
                if unit is None:
                    raise ProductImportError(
                        f"Baris {row_number}: satuan "
                        f"{normalized['unit']} ({normalized['unit_symbol']}) "
                        "belum tersedia."
                    )

                sku = str(normalized["sku"]).strip()
                product = self.session.scalar(
                    select(Product).where(
                        func.lower(Product.sku) == sku.casefold(),
                        Product.deleted_at.is_(None),
                    )
                )
                supplier_id = self._resolve_supplier(
                    str(normalized["supplier_code"]).strip(),
                    allow_missing=allow_missing_suppliers,
                )
                if product is None:
                    product = Product(
                        sku=sku,
                        barcode=str(normalized["barcode"]) or None,
                        name=str(normalized["name"]),
                        category_id=category.id,
                        unit_id=unit.id,
                        supplier_id=supplier_id,
                        purchase_price=int(normalized["purchase_price"]),
                        selling_price=int(normalized["selling_price"]),
                        stock_qty=int(normalized["stock_qty"]),
                        min_stock=int(normalized["min_stock"]),
                        is_active=bool(normalized["is_active"]),
                    )
                    self.session.add(product)
                    self.session.flush()
                    if product.stock_qty > 0:
                        self.session.add(
                            StockMovement(
                                product_id=product.id,
                                movement_type="INITIAL_STOCK",
                                stock_before=0,
                                qty=product.stock_qty,
                                stock_after=product.stock_qty,
                                reference_type="PRODUCT",
                                reference_id=product.id,
                                note="Stok awal dari impor produk.",
                            )
                        )
                    created += 1
                else:
                    product.barcode = str(normalized["barcode"]) or None
                    product.name = str(normalized["name"])
                    product.category_id = category.id
                    product.unit_id = unit.id
                    product.supplier_id = supplier_id
                    product.purchase_price = int(normalized["purchase_price"])
                    product.selling_price = int(normalized["selling_price"])
                    product.min_stock = int(normalized["min_stock"])
                    product.is_active = bool(normalized["is_active"])
                    updated += 1
            self.session.commit()
        except (AuthorizationError, ValidationError, SQLAlchemyError, ValueError) as exc:
            self.session.rollback()
            raise ProductImportError(str(exc)) from exc
        except Exception:
            self.session.rollback()
            raise

        return {"created": created, "updated": updated, "errors": errors}

    def export_template(self, file_path: str | Path) -> Path:
        path = Path(file_path)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Produk"
        sheet.append(PRODUCT_HEADERS)
        for product in self.master_data.list_products():
            category = self.session.get(Category, product.category_id)
            unit = self.session.get(Unit, product.unit_id)
            supplier = (
                self.session.get(Supplier, product.supplier_id)
                if product.supplier_id is not None
                else None
            )
            sheet.append(
                [
                    product.sku,
                    product.barcode or "",
                    product.name,
                    category.name if category else "",
                    unit.name if unit else "",
                    unit.symbol if unit else "",
                    supplier.code if supplier else "",
                    product.purchase_price,
                    product.selling_price,
                    product.stock_qty,
                    product.min_stock,
                    "Ya" if product.is_active else "Tidak",
                ]
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(
                max(len(str(cell.value or "")) for cell in column) + 2,
                32,
            )
            sheet.column_dimensions[column[0].column_letter].width = width
        workbook.save(path)
        return path

    def _read_csv(self, path: Path) -> list[dict[str, object]]:
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as file:
                reader = csv.DictReader(file)
                header_mapping = self._validate_headers(reader.fieldnames)
                return [
                    self._record_from_mapping(
                        {
                            header_mapping[index]: value
                            for index, value in enumerate(row.values())
                        }
                    )
                    for row in reader
                    if any(str(value or "").strip() for value in row.values())
                ]
        except (OSError, UnicodeError, csv.Error) as exc:
            raise ProductImportError(f"CSV tidak dapat dibaca: {exc}") from exc

    def _read_excel(self, path: Path) -> list[dict[str, object]]:
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            header_mapping = self._validate_headers(list(headers) if headers else None)
            records = []
            for values in rows:
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                records.append(
                    self._record_from_mapping(
                        {
                            header_mapping[index]: value
                            for index, value in enumerate(values)
                            if index < len(header_mapping)
                        }
                    )
                )
            return records
        except (OSError, ValueError, KeyError) as exc:
            raise ProductImportError(f"Excel tidak dapat dibaca: {exc}") from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _read_database(self, path: Path) -> list[dict[str, object]]:
        if not path.is_file():
            raise ProductImportError("File database tidak ditemukan.")
        connection = None
        try:
            database_uri = path.resolve().as_uri() + "?mode=ro"
            connection = sqlite3.connect(database_uri, uri=True)
            tables = {
                row[0]
                for row in connection.execute(
                    "SELECT name FROM sqlite_master WHERE type = 'table'"
                )
            }
            required = {"products", "categories", "units"}
            if not required.issubset(tables):
                raise ProductImportError(
                    "Database tidak kompatibel. Pilih file SQLite POS yang "
                    "memiliki tabel products, categories, dan units."
                )

            product_columns = {
                row[1] for row in connection.execute("PRAGMA table_info(products)")
            }
            required_columns = {
                "sku",
                "barcode",
                "name",
                "category_id",
                "unit_id",
                "purchase_price",
                "selling_price",
                "stock_qty",
                "min_stock",
                "is_active",
            }
            if not required_columns.issubset(product_columns):
                raise ProductImportError(
                    "Struktur tabel produk pada database tidak kompatibel "
                    "dengan aplikasi ini."
                )

            supplier_expr = (
                "supplier.code"
                if "suppliers" in tables and "supplier_id" in product_columns
                else "''"
            )
            supplier_join = (
                "LEFT JOIN suppliers AS supplier "
                "ON supplier.id = products.supplier_id"
                if supplier_expr != "''"
                else ""
            )
            deleted_filter = (
                "WHERE products.deleted_at IS NULL"
                if "deleted_at" in product_columns
                else ""
            )
            query = f"""
                SELECT products.sku, products.barcode, products.name,
                       categories.name, units.name, units.symbol,
                       {supplier_expr}, products.purchase_price,
                       products.selling_price, products.stock_qty,
                       products.min_stock, products.is_active
                FROM products
                LEFT JOIN categories ON categories.id = products.category_id
                LEFT JOIN units ON units.id = products.unit_id
                {supplier_join}
                {deleted_filter}
                ORDER BY products.sku
            """
            records = []
            for values in connection.execute(query):
                record = self._record_from_mapping(
                    dict(zip(PRODUCT_HEADERS, values))
                )
                records.append(record)
            return records
        except sqlite3.Error as exc:
            raise ProductImportError(
                f"Database SQLite tidak dapat dibaca: {exc}"
            ) from exc
        finally:
            if connection is not None:
                connection.close()

    @staticmethod
    def _validate_headers(headers: list[str | None] | None) -> list[str]:
        if headers is None:
            raise ProductImportError("File tidak memiliki baris judul kolom.")
        normalized = [str(header or "").strip().casefold() for header in headers]
        required = {header.casefold() for header in PRODUCT_HEADERS}
        if len(normalized) != len(PRODUCT_HEADERS) or set(normalized) != required:
            raise ProductImportError(
                "Judul kolom tidak sesuai template. Gunakan file template "
                "Excel dari menu Edit Excel."
            )
        canonical = {header.casefold(): header for header in PRODUCT_HEADERS}
        return [canonical[header] for header in normalized]

    @staticmethod
    def _record_from_mapping(row: dict[str, object]) -> dict[str, object]:
        return {
            header: row.get(header, "")
            for header in PRODUCT_HEADERS
        }

    @staticmethod
    def _normalize_record(row: dict[str, object]) -> dict[str, object]:
        def required_text(header: str) -> str:
            value = str(row.get(header) or "").strip()
            if not value:
                raise ProductImportError(f"Kolom {header} wajib diisi.")
            return value

        def number(header: str) -> int:
            value = row.get(header)
            try:
                if isinstance(value, (int, float)):
                    decimal_value = Decimal(str(value))
                else:
                    normalized = str(value or "0").strip()
                    decimal_value = Decimal(
                        normalized.replace(".", "").replace(",", "")
                    )
                numeric = int(decimal_value)
                if decimal_value != numeric:
                    raise ValueError
            except (InvalidOperation, TypeError, ValueError) as exc:
                raise ProductImportError(
                    f"Kolom {header} harus berupa bilangan bulat tanpa pecahan."
                ) from exc
            if numeric < 0:
                raise ProductImportError(f"Kolom {header} tidak boleh negatif.")
            return numeric

        active = str(row.get("Aktif") or "Ya").strip().casefold()
        if active in {"ya", "yes", "true", "1", "aktif"}:
            is_active = True
        elif active in {"tidak", "no", "false", "0", "nonaktif"}:
            is_active = False
        else:
            raise ProductImportError("Kolom Aktif hanya boleh Ya atau Tidak.")
        return {
            "sku": required_text("SKU"),
            "barcode": str(row.get("Barcode") or "").strip(),
            "name": required_text("Nama Produk"),
            "category": required_text("Kategori"),
            "unit": required_text("Satuan"),
            "unit_symbol": required_text("Simbol Satuan"),
            "supplier_code": str(row.get("Kode Supplier") or "").strip(),
            "purchase_price": number("Harga Modal"),
            "selling_price": number("Harga Jual"),
            "stock_qty": number("Stok Awal"),
            "min_stock": number("Batas Minimum Stok"),
            "is_active": is_active,
        }

    def _find_category(self, name: str) -> Category | None:
        return self.session.scalar(
            select(Category).where(func.lower(Category.name) == name.casefold())
        )

    def _find_unit(self, name: str, symbol: str) -> Unit | None:
        return self.session.scalar(
            select(Unit).where(
                func.lower(Unit.name) == name.casefold(),
                func.lower(Unit.symbol) == symbol.casefold(),
            )
        )

    def _resolve_supplier(
        self,
        code: str,
        *,
        allow_missing: bool = False,
    ) -> int | None:
        if not code:
            return None
        supplier = self.session.scalar(
            select(Supplier).where(
                func.lower(Supplier.code) == code.casefold(),
                Supplier.deleted_at.is_(None),
            )
        )
        if supplier is None:
            if allow_missing:
                return None
            raise ProductImportError(
                f"Supplier dengan kode {code} tidak ditemukan."
            )
        return supplier.id

    def _validate_references(
        self,
        row: dict[str, object],
        *,
        create_missing_units: bool,
        allow_missing_suppliers: bool,
    ):
        category_name = str(row["category"]).strip()
        category = self._find_category(category_name)
        if category is None:
            existing = self.session.scalar(
                select(Category).where(
                    func.lower(Category.name) == category_name.casefold()
                )
            )
            if existing is not None and existing.deleted_at is None:
                raise ProductImportError(
                    f"Kategori {category_name} tidak dapat digunakan."
                )
        unit_name = str(row["unit"]).strip()
        unit_symbol = str(row["unit_symbol"]).strip()
        unit = self._find_unit(unit_name, unit_symbol)
        if unit is None:
            same_name_unit = self.session.scalar(
                select(Unit).where(func.lower(Unit.name) == unit_name.casefold())
            )
            if same_name_unit is not None:
                raise ProductImportError(
                    f"Satuan {unit_name} sudah ada dengan simbol "
                    f"{same_name_unit.symbol}, bukan {unit_symbol}."
                )
        if unit is None and not create_missing_units:
            raise ProductImportError(
                f"Satuan {unit_name} ({unit_symbol}) belum tersedia."
            )
        self._resolve_supplier(
            str(row["supplier_code"]).strip(),
            allow_missing=allow_missing_suppliers,
        )
